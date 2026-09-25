# -*- coding: utf-8 -*-
"""V123 — Bộ câu hỏi ĐẦY ĐỦ CHI TIẾT THEO TỔ & CÔNG ĐOẠN (module Lên kế hoạch sản xuất) — phần 1/2."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re as _re

DO = RGBColor(0x0E, 0x74, 0x90); XAM = RGBColor(0x55, 0x55, 0x55); RED = RGBColor(0xB9, 0x1C, 0x1C)
XANH = RGBColor(0x0B, 0x4A, 0x6F)
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.left_margin = s.right_margin = Cm(1.1); s.top_margin = Cm(1.1); s.bottom_margin = Cm(1.2)
n = doc.styles["Normal"]; n.font.name = "Arial"; n.font.size = Pt(9)
n._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial"); n._element.rPr.rFonts.set(qn("w:cs"), "Arial")
WT = Cm(21 - 2.2)


def run(p, t, size=9, bold=False, italic=False, color=None):
    r = p.add_run(t); r.font.name = "Arial"; r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color is not None: r.font.color.rgb = color
    r._element.rPr.rFonts.set(qn("w:cs"), "Arial"); return r


def rich(p, text, size=9, bold=False, italic=False, color=None):
    for part in _re.split(r'(<b>.*?</b>)', text):
        if not part: continue
        if part.startswith('<b>') and part.endswith('</b>'):
            run(p, part[3:-4], size=size, bold=True, italic=italic, color=color)
        else:
            run(p, part, size=size, bold=bold, italic=italic, color=color)


def shade(cell, c):
    el = OxmlElement("w:shd"); el.set(qn("w:val"), "clear"); el.set(qn("w:fill"), c)
    cell._tc.get_or_add_tcPr().append(el)


def head(t, size=11.5):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(3)
    pPr = p._p.get_or_add_pPr()
    el = OxmlElement("w:shd"); el.set(qn("w:val"), "clear"); el.set(qn("w:fill"), "DDEEF5"); pPr.append(el)
    bd = OxmlElement("w:pBdr"); l = OxmlElement("w:left")
    l.set(qn("w:val"), "single"); l.set(qn("w:sz"), "20"); l.set(qn("w:color"), "0E7490")
    bd.append(l); pPr.append(bd)
    run(p, " " + t, size=size, bold=True, color=RGBColor(0x08, 0x33, 0x44))


def sub(t):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
    rich(p, t, size=9.5, bold=True, color=XANH)


def note(t):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    rich(p, t, size=8, italic=True, color=XAM)


def qs(rows):
    ws = [Cm(1.3), WT - Cm(1.3) - Cm(1.2) - Cm(5.7), Cm(1.2), Cm(5.7)]
    table = doc.add_table(rows=1, cols=4); table.style = "Table Grid"
    for i, t in enumerate(["#", "Câu hỏi — vì sao cần hỏi", "Mức", "Trả lời"]):
        c = table.rows[0].cells[i]; shade(c, "EEEEEE")
        p = c.paragraphs[0]
        if i in (0, 2): p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p, t, size=8, bold=True)
    trPr = table.rows[0]._tr.get_or_add_trPr(); trPr.append(OxmlElement("w:tblHeader"))
    for code, q, why, lv in rows:
        cells = table.add_row().cells
        table.rows[-1].height = Cm(0.5); table.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        p = cells[0].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p, code, size=8, bold=True); shade(cells[0], "F7F7F7")
        p = cells[1].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        rich(p, q, size=8.5, bold=True)
        if why:
            pw = cells[1].add_paragraph(); pw.paragraph_format.space_before = Pt(0)
            run(pw, why, size=7, italic=True, color=XAM)
        p = cells[2].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p, lv, size=8, bold=True, color=RED if lv == "P1" else None)
        shade(cells[3], "FFFDF2")
    table.autofit = False
    for r in table.rows:
        for i, w in enumerate(ws): r.cells[i].width = w
    return table


def tbl(header, rows, ws, tick=True):
    table = doc.add_table(rows=1, cols=len(ws)); table.style = "Table Grid"
    for i, t in enumerate(header):
        c = table.rows[0].cells[i]; shade(c, "EEEEEE")
        run(c.paragraphs[0], t, size=8, bold=True)
    trPr = table.rows[0]._tr.get_or_add_trPr(); trPr.append(OxmlElement("w:tblHeader"))
    for r in rows:
        cells = table.add_row().cells
        table.rows[-1].height = Cm(0.5); table.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for i, v in enumerate(r):
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            rich(p, str(v), size=8.5, bold=(i == 0))
        if tick: shade(cells[len(r) - 1], "FFFDF2")
    table.autofit = False
    for row in table.rows:
        for i, w in enumerate(ws): row.cells[i].width = w
    return table


# ================= TIÊU ĐỀ =================
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
run(p, "BỘ CÂU HỎI ĐẦY ĐỦ THEO TỔ & CÔNG ĐOẠN — MODULE LÊN KẾ HOẠCH SẢN XUẤT", size=13.5, bold=True)
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5)
run(p, "Công ty: ", size=8, bold=True); run(p, ".....................................     ", size=8)
run(p, "Ngày: ", size=8, bold=True); run(p, "....../....../2026     ", size=8)
run(p, "Người trả lời: ", size=8, bold=True); run(p, "..................................... (quản đốc / tổ trưởng / kỹ thuật)\n", size=8)
run(p, "Cách dùng: mỗi tổ trả lời phần của mình (mục 1 → 6), người lên kế hoạch trả lời mục 7. Mức ", size=8)
run(p, "P1", size=8, bold=True, color=RED)
run(p, " = bắt buộc để lập trình; ", size=8); run(p, "P2", size=8, bold=True)
run(p, " = làm sau cũng được. Chưa rõ thì ghi “chưa rõ” — app chạy theo mặc định ở mục 10.\n", size=8)
run(p, "Trong mỗi công đoạn, câu hỏi đi theo đúng 7 nhóm để tổ dễ trả lời: ", size=8, italic=True, color=XAM)
run(p, "① nội dung & phạm vi · ② thời lượng · ③ nguồn lực & máy móc · ④ setup & gom lô · ⑤ chờ & ràng buộc · ⑥ kiểm tra & lỗi · ⑦ ghi nhận trong app.", size=8, italic=True, color=XAM)

# ================= 1. TỔ MÁY =================
head("1. TỔ MÁY  (Thiết kế · Bồi Lares · Cắt · Chấn)")
sub("1.1. Công đoạn THIẾT KẾ")
qs([
 ("TK1", "① 1 bộ cửa cần những bản vẽ gì (cắt nhôm, cánh/panel, kính, lắp ráp)? Có bản vẽ khung bao / ô thoáng riêng không?", "app cần biết phải lưu/ in gì theo bộ", "P1"),
 ("TK2", "② Một bộ (hoặc một model) thiết kế mất bao lâu? Có phụ thuộc loại cửa / kích thước không?", "định mức thời gian của bước chuẩn bị", "P1"),
 ("TK3", "③ Ai làm, mấy người? Dùng phần mềm gì? Làm ở văn phòng hay tại xưởng?", "đây là nguồn lực riêng, không cộng vào năng lực tổ thợ", "P1"),
 ("TK4", "④ Model/mã hàng đã từng làm thì có dùng lại bản vẽ cũ không? Khách đổi thì sửa mất bao lâu?", "nếu tái sử dụng được thì thời lượng = 0 với mã quen", "P1"),
 ("TK5", "⑤ Thiết kế cần những thông tin gì từ đơn (cao/rộng · khung bao · thông thủy · số cánh · ô thoáng · màu · loại kính · hướng mở)? App đang có đủ chưa, còn thiếu gì?", "đây là danh sách trường app phải bổ sung/bắt buộc", "P1"),
 ("TK6", "⑥ Ai duyệt bản vẽ? Khách có phải duyệt trước khi sản xuất không? Duyệt xong mới được cắt đúng không?", "mốc “đủ điều kiện vào sản xuất”", "P1"),
 ("TK7", "⑦ Có cần lưu bản vẽ lên app theo từng bộ để tổ xem tại xưởng không?", "thay việc gửi Zalo thủ công", "P2"),
])
sub("1.2. Công đoạn <b>BỒI LARES</b> — nạp chương trình cho MÁY CẮT CHUYÊN DỤNG (hậu thiết kế)")
note("Đã xác nhận: Bồi Lares là bước chuyển file thiết kế sang chương trình cho máy cắt chuyên dụng (không phải vật liệu, không phải ép cốt).")
qs([
 ("BL1", "① Chương trình làm theo <b>MODEL NHÔM</b> hay theo <b>TỪNG BỘ / kích thước</b>? Một chương trình dùng lại được cho các đơn sau không?", "nếu theo model thì 1 chương trình dùng cho nhiều bộ — tải tính theo model", "P1"),
 ("BL2", "③ Ai làm chương trình (kỹ thuật / văn phòng hay tại máy cắt)? Mấy người? Dùng phần mềm gì?", "xác định đây là nguồn lực riêng (nhóm KỸ THUẬT/CAM)", "P1"),
 ("BL3", "② Chương trình <b>MỚI</b> (model chưa từng làm) mất bao lâu? Chương trình <b>ĐÃ CÓ</b> thì nạp lại/sửa mất bao lâu?", "hai con số khác nhau — app cần cả hai để tính tải", "P1"),
 ("BL4", "⑤ Chưa có chương trình thì máy cắt có chạy được không? Đã bao giờ <b>máy cắt đứng chờ chương trình</b> chưa?", "đây là ràng buộc bắt buộc (prerequisite) của công đoạn Cắt", "P1"),
 ("BL5", "① Một chương trình gồm những phần nào (khung / cánh / phào)? Là 1 file cho cả bộ hay mỗi phần 1 file?", "biết đơn vị theo dõi: theo bộ hay theo bộ phận", "P2"),
 ("BL6", "④ Chương trình lưu ở đâu (phần mềm máy / máy tính / USB)? Đặt tên theo gì? Khách đổi thì tạo phiên bản mới thế nào?", "để dựng thư viện chương trình theo model trong app", "P2"),
 ("BL7", "③ Xưởng có mấy <b>máy cắt chuyên dụng</b>? Model máy gì, đọc định dạng file nào, có phải chuyển file bằng tay (USB) không?", "năng lực cắt thật + chương trình phải khớp máy", "P2"),
 ("BL8", "⑥ Có <b>chạy thử / cắt mẫu</b> 1 bộ trước khi cắt hàng loạt không? Mất bao lâu? Có phải kiểm mẫu rồi mới chạy tiếp không?", "nếu có thì thêm một bước và một khoản thời gian trước khi cắt", "P2"),
 ("BL9", "⑤ Khách đổi kích thước/mẫu giữa lúc đang làm thì <b>sửa chương trình</b> mất bao lâu? Có phải cắt lại từ đầu không?", "thời gian phát sinh khi đơn thay đổi giữa chừng", "P1"),
])
sub("1.3. Công đoạn CẮT")
qs([
 ("C1", "③ Cắt bằng máy gì (máy chuyên dụng / CNC / cắt tay)? Mấy máy, mấy người đứng máy?", "năng lực thật của công đoạn cắt", "P1"),
 ("C2", "① Cắt cho những phần nào: khung · cánh · phào? Là 3 việc làm song song hay 1 người làm lần lượt?", "3 luồng song song — ảnh hưởng cách xếp tải", "P1"),
 ("C3", "② Thời gian cắt tính theo gì: số bộ · số mét dài nhôm · số thanh? Một bộ mất bao lâu?", "công thức định mức thời gian", "P1"),
 ("C4", "④ Đổi sang model nhôm khác mất bao lâu (đổi dao, cài đặt, lấy nhôm)? Có gom các bộ cùng model để cắt chung không?", "thời gian setup — quyết định có nên gom lô cắt", "P1"),
 ("C5", "⑤ Nhôm có sẵn kho theo model hay phải đặt theo đơn? Chờ bao lâu? App có cần cảnh báo “thiếu nhôm” trước khi xếp lịch?", "nếu chờ lâu thì kế hoạch phải chừa thời gian vật tư", "P1"),
 ("C6", "⑤ Cắt có phải chờ <b>chương trình Bồi Lares</b> không? Thường chờ bao lâu?", "đo mức độ rủi ro của ràng buộc bắt buộc", "P1"),
 ("C7", "⑥ Sau cắt có kiểm tra kích thước / số lượng không? Ai kiểm? Cắt sai/hụt xử lý thế nào, mất bao lâu?", "điểm QC + xử lý lỗi", "P2"),
 ("C8", "③② Một ngày cắt được bao nhiêu bộ? Nhôm cắt xong để chờ chấn được bao lâu, có giới hạn chỗ không?", "năng lực + giới hạn hàng chờ", "P1"),
])
sub("1.4. Công đoạn CHẤN")
qs([
 ("CH1", "③ Có mấy máy chấn? Cần khuôn riêng theo model nhôm không? Hiện có mấy khuôn?", "nút thắt theo khuôn, không chỉ theo người", "P1"),
 ("CH2", "① Chấn khung / cánh / phào trên cùng 1 máy hay máy khác nhau? Có bị tranh máy giữa 3 luồng không?", "tranh chấp máy giữa 3 luồng khung/cánh/phào", "P1"),
 ("CH3", "② Một bộ chấn mất bao lâu? Có phụ thuộc kích thước cửa / độ dài thanh không?", "định mức thời gian", "P1"),
 ("CH4", "④ Đổi khuôn / đổi model mất bao lâu? Có gom lô theo model để đỡ đổi khuôn không? Tối thiểu bao nhiêu bộ 1 lượt?", "thời gian setup + quy tắc gom lô", "P1"),
 ("CH5", "⑤ Ghi chú trong bảng: <b>“cần đầy đủ đơn để đẩy công đoạn kế tiếp”</b> — nghĩa là đủ <b>3 phần của 1 bộ</b> (khung+cánh+phào), hay đủ <b>tất cả các bộ trong đơn</b>?", "quyết định app cho đẩy từng bộ hay bắt buộc theo đơn — câu chặn quan trọng nhất", "P1"),
 ("CH6", "⑥ Sau chấn có kiểm tra góc / kích thước không? Ai kiểm? Chấn lỗi (méo, lệch góc) sửa được không hay phải cắt lại?", "điểm QC + xử lý lỗi", "P2"),
 ("CH7", "③② Một ngày chấn được bao nhiêu bộ? Hàng đã chấn để chờ hàn ở đâu, được bao lâu?", "năng lực + chỗ để hàng chờ", "P1"),
])
sub("1.5. TỔ MÁY — câu hỏi chung")
qs([
 ("TM1", "③ Tổ máy bao nhiêu người? Chia thế nào cho 4 việc: thiết kế · Bồi Lares · cắt · chấn?", "năng lực theo người", "P1"),
 ("TM2", "③ Làm mấy ca? Giờ bắt đầu – kết thúc? Nghỉ trưa bao lâu?", "quy đổi ngày công ra giờ làm việc", "P1"),
 ("TM3", "② Một ngày tổ máy ra được bao nhiêu bộ “sạch” (đã chấn xong, sẵn sàng chuyển hàn)?", "đầu ra của tổ để tính tải", "P1"),
 ("TM4", "⑥ Máy nào hay hỏng / phải bảo trì? Bao lâu bảo trì 1 lần, mất bao lâu?", "app phải chừa thời gian dừng máy, không xếp kín", "P2"),
 ("TM5", "⑤ Việc nào ở tổ máy hay bị kẹt nhất: chờ nhôm · chờ chương trình · chờ thiết kế · chờ máy?", "nguyên nhân trễ để app cảnh báo đúng chỗ", "P1"),
 ("TM6", "⑦ Tổ máy cập nhật tiến độ ở đâu? Ai cập nhật (tổ trưởng hay văn phòng)?", "quyết định giao diện nhập liệu", "P1"),
])

doc.add_page_break()
# ================= 2. TỔ HÀN =================
head("2. TỔ HÀN  (Hàn · Ép cánh · Test cơ khí)")
sub("2.1. Công đoạn HÀN")
qs([
 ("HAN1", "③ Hàn bằng máy gì, mấy máy, mấy người?", "năng lực", "P1"),
 ("HAN2", "① Hàn những phần nào: khung · cánh · phào? 3 người làm song song hay 1 người tuần tự?", "3 luồng song song hay tuần tự", "P1"),
 ("HAN3", "② Một bộ hàn mất bao lâu? Tính theo số góc / mối hàn hay theo bộ?", "định mức thời gian", "P1"),
 ("HAN4", "④ Có phải gom lô theo model để đỡ đổi khuôn/gá không? Tối thiểu bao nhiêu bộ 1 lượt?", "quy tắc gom lô hàn", "P1"),
 ("HAN5", "⑤ Có phải chờ <b>đủ 3 phần</b> (khung, cánh, phào) mới hàn được không? Hàng chờ ở đâu, tối đa bao nhiêu bộ?", "điều kiện vào công đoạn + giới hạn hàng chờ", "P1"),
 ("HAN6", "⑥ Sau hàn có <b>mài mối hàn / làm sạch ba via</b> không? Ai làm, mất bao lâu? Có kiểm mối hàn không?", "bước hoàn thiện thường bị bỏ sót khi định mức", "P1"),
 ("HAN7", "⑥⑦ Mối hàn lỗi xử lý thế nào, mất bao lâu? Một ngày tổ hàn được bao nhiêu bộ?", "xử lý lỗi + năng lực", "P1"),
])
sub("2.2. Công đoạn ÉP CÁNH  (bảng công đoạn đang để trống thời gian)")
qs([
 ("EP1", "① Ép cánh là ép cái gì vào cánh (tấm cốt / panel nhôm / kính)? Làm cho loại cửa nào?", "xác định đúng công đoạn và phạm vi áp dụng", "P1"),
 ("EP2", "③ Có máy ép không, mấy cái? Mỗi lượt ép được bao nhiêu cánh?", "năng lực + gom lô ép", "P1"),
 ("EP3", "② Một bộ (hoặc 1 cánh) ép mất bao lâu? Có phải <b>chờ keo khô</b> không, bao lâu?", "định mức + thời gian chờ", "P1"),
 ("EP4", "⑤ Ép cánh nằm ở đâu trong chuỗi: trước hay sau khi hàn cánh? Trước hay sau Test cơ khí?", "thứ tự công đoạn trong kế hoạch", "P1"),
 ("EP5", "④⑤ Vật tư (keo, cốt) lấy từ kho hay đặt ngoài? Có bị thiếu làm dừng việc không?", "ảnh hưởng lead time và cảnh báo vật tư", "P2"),
 ("EP6", "⑥ Ép lỗi (bọt khí, lệch, bong) xử lý thế nào? Có phải làm lại từ đầu không? Mất bao lâu?", "làm lại — ảnh hưởng lớn tới lead time", "P2"),
 ("EP7", "②③ Một ngày ép được bao nhiêu bộ?", "năng lực", "P1"),
])
sub("2.3. Công đoạn TEST CƠ KHÍ")
qs([
 ("TC1", "① Test gồm những gì (đóng/mở, khe hở cánh, khoá, bản lề, gioăng, độ vuông, kín nước)?", "app có thể hiện checklist QC — cần biết nội dung", "P1"),
 ("TC2", "②③ Ai test, mất bao lâu 1 bộ? Test từng bộ hay chỉ test mẫu?", "thời lượng + tần suất", "P1"),
 ("TC3", "⑥⑦ Kết quả test ghi ở đâu (giấy/app)? Có cần lưu lịch sử theo bộ trong app không? Đạt thì có dán tem/đánh dấu không?", "app có cần form nhập kết quả QC", "P2"),
 ("TC4", "⑥ Tỷ lệ phải sửa sau test khoảng bao nhiêu %? Lỗi thường gặp là gì?", "để dự phòng thời gian làm lại trong kế hoạch", "P1"),
 ("TC5", "⑤⑥ Phát hiện lỗi ở test thì quay lại công đoạn nào (hàn? chấn? ép cánh?)? Việc làm lại có chen trước đơn thường không?", "xác định luồng quay lại khi làm lại", "P1"),
 ("TC6", "②③ Một ngày test được bao nhiêu bộ?", "năng lực", "P2"),
])
sub("2.4. TỔ HÀN — câu hỏi chung")
qs([
 ("TH1", "③ Tổ hàn bao nhiêu người, mấy ca, giờ làm?", "năng lực theo người/giờ", "P1"),
 ("TH2", "⑤ Hàng sau khi test cơ khí <b>chờ sơn</b> bao lâu? Chỗ để có giới hạn không (tối đa bao nhiêu bộ)?", "giới hạn hàng chờ trước sơn — thường là nút thắt", "P1"),
 ("TH3", "② Một ngày tổ hàn ra được bao nhiêu bộ hoàn chỉnh (đã test đạt)?", "năng lực thực tế", "P1"),
 ("TH4", "⑥ Máy hàn/máy ép hay hỏng không? Bảo trì bao lâu 1 lần?", "thời gian dừng máy", "P2"),
 ("TH5", "⑤ Tổ hàn có bị ảnh hưởng khi tổ máy giao thiếu hoặc không đúng model không? Có phải chờ không?", "rủi ro phụ thuộc giữa các tổ", "P1"),
])
